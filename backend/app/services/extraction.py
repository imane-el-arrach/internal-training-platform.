from functools import lru_cache
from pathlib import Path
import re

from pptx import Presentation
from pypdf import PdfReader

EXTENSIONS_VIDEO = {".mp4", ".mov", ".avi", ".mkv", ".webm"}

# PostgreSQL refuse les caractères NUL dans les colonnes texte. Certains PDF
# en contiennent après extraction (polices ou données internes), alors qu'ils
# ne sont pas visibles dans le document. Les autres caractères de contrôle,
# hors espaces et retours à la ligne, sont également écartés avant le découpage
# et la génération des embeddings.
_CARACTERES_CONTROLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def nettoyer_texte_indexable(texte: str) -> str:
    texte_sans_controles = _CARACTERES_CONTROLE.sub(" ", texte)
    return re.sub(r"[ \t]+", " ", texte_sans_controles).strip()


def extraire_texte_pdf(chemin: str) -> str:
    lecteur = PdfReader(chemin)
    pages = [page.extract_text() or "" for page in lecteur.pages]
    return "\n\n".join(p.strip() for p in pages if p.strip())


def extraire_texte_pptx(chemin: str) -> str:
    presentation = Presentation(chemin)
    textes_slides = []

    for slide in presentation.slides:
        morceaux = []
        for forme in slide.shapes:
            if forme.has_text_frame:
                texte_forme = forme.text_frame.text.strip()
                if texte_forme:
                    morceaux.append(texte_forme)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            morceaux.append(slide.notes_slide.notes_text_frame.text.strip())

        if morceaux:
            textes_slides.append("\n".join(morceaux))

    return "\n\n".join(textes_slides)


@lru_cache(maxsize=1)
def _get_modele_whisper():
    from faster_whisper import WhisperModel

    from app.core.config import settings

    return WhisperModel(
        settings.WHISPER_MODEL,
        device=settings.WHISPER_DEVICE,
        compute_type=settings.WHISPER_COMPUTE_TYPE,
    )


def extraire_texte_video(chemin: str) -> str:
    """
    Transcription audio -> texte. `language=None` laisse Whisper détecter
    la langue automatiquement (nos formations peuvent être en français ou
    en anglais) 
    """
    modele = _get_modele_whisper()
    segments, _info = modele.transcribe(chemin, language=None)
    return " ".join(segment.text.strip() for segment in segments)


def extraire_texte(chemin: str) -> str:
    extension = Path(chemin).suffix.lower()

    if extension == ".pdf":
        return nettoyer_texte_indexable(extraire_texte_pdf(chemin))
    if extension in (".pptx", ".ppt"):
        return nettoyer_texte_indexable(extraire_texte_pptx(chemin))
    if extension in EXTENSIONS_VIDEO:
        return nettoyer_texte_indexable(extraire_texte_video(chemin))

    raise ValueError(
        f"Extraction de texte non supportée pour l'extension '{extension}'"
    )

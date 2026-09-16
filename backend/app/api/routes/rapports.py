import csv
import io
import uuid
from datetime import date, datetime

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from app.api.deps import exiger_administrateur
from app.db.session import get_session
from app.models.affectation import Affectation
from app.models.departement import Departement
from app.models.formation import Formation
from app.models.progression import Progression, StatutProgression
from app.models.questionnaire import Questionnaire
from app.models.tentative import StatutTentative, Tentative
from app.models.utilisateur import Utilisateur
from app.schemas.rapport import RapportSuiviRead, ResultatIndividuelRead
from app.services.rapport_pdf import generer_pdf_rapport

router = APIRouter(prefix="/api/rapports", tags=["Rapports"])


@router.get("/suivi", response_model=RapportSuiviRead)
def consulter_suivi(
    formation_id: uuid.UUID | None = Query(default=None),
    departement_id: uuid.UUID | None = Query(default=None),
    _: Utilisateur = Depends(exiger_administrateur),
    session: Session = Depends(get_session),
) -> RapportSuiviRead:
    """Synthèse de pilotage, servant aussi de base aux futurs exports."""
    formations = {formation.id: formation for formation in session.exec(select(Formation)).all()}
    utilisateurs = {
        utilisateur.id: utilisateur for utilisateur in session.exec(select(Utilisateur)).all()
    }
    departements = {
        departement.id: departement
        for departement in session.exec(select(Departement)).all()
    }
    affectations = {
        affectation.id: affectation
        for affectation in session.exec(select(Affectation)).all()
    }

    progressions = list(session.exec(select(Progression)).all())
    if formation_id is not None:
        progressions = [p for p in progressions if p.formation_id == formation_id]
    if departement_id is not None:
        progressions = [
            p
            for p in progressions
            if utilisateurs.get(p.utilisateur_id)
            and utilisateurs[p.utilisateur_id].departement_id == departement_id
        ]

    resultats: list[ResultatIndividuelRead] = []
    for progression in progressions:
        utilisateur = utilisateurs.get(progression.utilisateur_id)
        formation = formations.get(progression.formation_id)
        if utilisateur is None or formation is None:
            continue

        departement = departements.get(utilisateur.departement_id)
        affectation = affectations.get(progression.affectation_id)
        resultats.append(
            ResultatIndividuelRead(
                utilisateur_id=utilisateur.id,
                utilisateur_nom=utilisateur.nom,
                utilisateur_prenom=utilisateur.prenom,
                departement_nom=departement.nom if departement else None,
                formation_id=formation.id,
                formation_titre=formation.titre,
                formation_obligatoire=formation.obligatoire,
                statut=progression.statut.value,
                pourcentage=progression.pourcentage,
                date_limite=affectation.date_limite if affectation else None,
                derniere_activite=progression.derniere_activite,
            )
        )

    total = len(resultats)
    terminees = sum(resultat.statut == StatutProgression.termine.value for resultat in resultats)
    en_attente = sum(
        resultat.statut == StatutProgression.non_commence.value for resultat in resultats
    )
    en_cours = sum(
        resultat.statut == StatutProgression.en_cours.value for resultat in resultats
    )
    taux_completion = round(sum(resultat.pourcentage for resultat in resultats) / total) if total else 0

    obligatoires = [resultat for resultat in resultats if resultat.formation_obligatoire]
    taux_conformite = (
        round(sum(resultat.statut == StatutProgression.termine.value for resultat in obligatoires) * 100 / len(obligatoires))
        if obligatoires
        else None
    )

    formation_ids = {resultat.formation_id for resultat in resultats}
    utilisateur_ids = {resultat.utilisateur_id for resultat in resultats}
    questionnaires = {
        questionnaire.id: questionnaire
        for questionnaire in session.exec(select(Questionnaire)).all()
        if questionnaire.formation_id in formation_ids
    }
    tentatives_formation = [
        tentative
        for tentative in session.exec(select(Tentative)).all()
        if tentative.utilisateur_id in utilisateur_ids
        and tentative.questionnaire_id in questionnaires
    ]
    tentatives = [
        tentative
        for tentative in tentatives_formation
        if tentative.statut == StatutTentative.TERMINEE
    ]

    # Les anciennes progressions ne portent pas toujours une activité ; dans
    # ce cas, l'historique du quiz fournit une date fiable sans modifier le passé.
    activites_quiz: dict[tuple[uuid.UUID, uuid.UUID], datetime] = {}
    for tentative in tentatives_formation:
        formation_quiz_id = questionnaires[tentative.questionnaire_id].formation_id
        cle = (tentative.utilisateur_id, formation_quiz_id)
        activite = tentative.date_passage or tentative.date_debut
        precedente = activites_quiz.get(cle)
        if precedente is None or activite > precedente:
            activites_quiz[cle] = activite

    for resultat in resultats:
        activite_quiz = activites_quiz.get((resultat.utilisateur_id, resultat.formation_id))
        if activite_quiz is not None and (
            resultat.derniere_activite is None or activite_quiz > resultat.derniere_activite
        ):
            resultat.derniere_activite = activite_quiz

    # Une seule tentative est prise en compte par questionnaire : la dernière.
    dernieres_tentatives: dict[tuple[uuid.UUID, uuid.UUID], Tentative] = {}
    for tentative in tentatives:
        cle = (tentative.utilisateur_id, tentative.questionnaire_id)
        precedente = dernieres_tentatives.get(cle)
        if precedente is None or tentative.numero_tentative > precedente.numero_tentative:
            dernieres_tentatives[cle] = tentative

    valeurs_reussite = [
        tentative.reussi for tentative in dernieres_tentatives.values() if tentative.reussi is not None
    ]
    taux_reussite = (
        round(sum(valeurs_reussite) * 100 / len(valeurs_reussite))
        if valeurs_reussite
        else None
    )

    return RapportSuiviRead(
        collaborateurs_concernes=len({resultat.utilisateur_id for resultat in resultats}),
        progressions_total=total,
        progressions_terminees=terminees,
        formations_en_attente=en_attente,
        progressions_en_cours=en_cours,
        taux_completion=taux_completion,
        taux_reussite=taux_reussite,
        taux_conformite=taux_conformite,
        resultats=sorted(
            resultats,
            key=lambda resultat: (resultat.utilisateur_nom.lower(), resultat.formation_titre.lower()),
        ),
    )


@router.get("/suivi/export/csv")
def exporter_suivi_csv(
    formation_id: uuid.UUID | None = Query(default=None),
    departement_id: uuid.UUID | None = Query(default=None),
    admin: Utilisateur = Depends(exiger_administrateur),
    session: Session = Depends(get_session),
) -> StreamingResponse:
    """Exporte les mêmes résultats que la page de suivi, avec ses filtres."""
    rapport = consulter_suivi(
        formation_id=formation_id,
        departement_id=departement_id,
        _=admin,
        session=session,
    )

    contenu = io.StringIO(newline="")
    contenu.write("\ufeff")  # BOM UTF-8 : accents correctement lus par Excel.
    ecrivain = csv.writer(contenu, delimiter=";")
    ecrivain.writerow(
        [
            "Collaborateur",
            "Service",
            "Formation",
            "Obligatoire",
            "Progression",
            "Statut",
            "Échéance",
            "Dernière activité",
        ]
    )

    libelles_statut = {
        "non_commence": "En attente",
        "en_cours": "En cours",
        "termine": "Terminée",
    }
    for resultat in rapport.resultats:
        ecrivain.writerow(
            [
                f"{resultat.utilisateur_prenom} {resultat.utilisateur_nom}",
                resultat.departement_nom or "Non renseigné",
                resultat.formation_titre,
                "Oui" if resultat.formation_obligatoire else "Non",
                f"{resultat.pourcentage}%",
                libelles_statut[resultat.statut],
                resultat.date_limite.strftime("%d/%m/%Y") if resultat.date_limite else "",
                resultat.derniere_activite.strftime("%d/%m/%Y %H:%M")
                if resultat.derniere_activite
                else "Aucune activité",
            ]
        )

    nom_fichier = f"exia-academy-suivi-{date.today().isoformat()}.csv"
    return StreamingResponse(
        iter([contenu.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{nom_fichier}"'},
    )


@router.get("/suivi/export/excel")
def exporter_suivi_excel(
    formation_id: uuid.UUID | None = Query(default=None),
    departement_id: uuid.UUID | None = Query(default=None),
    admin: Utilisateur = Depends(exiger_administrateur),
    session: Session = Depends(get_session),
) -> StreamingResponse:
    """Génère un classeur Excel avec la synthèse et le détail filtré."""
    rapport = consulter_suivi(
        formation_id=formation_id,
        departement_id=departement_id,
        _=admin,
        session=session,
    )

    classeur = Workbook()
    synthese = classeur.active
    synthese.title = "Synthèse"
    detail = classeur.create_sheet("Détail individuel")

    bleu_fonce = "315864"
    turquoise = "49BFA9"
    entete = PatternFill("solid", fgColor=bleu_fonce)
    accent = PatternFill("solid", fgColor="E7F0EF")
    police_entete = Font(color="FFFFFF", bold=True)

    synthese.merge_cells("A1:B1")
    synthese["A1"] = "EXIA Academy — Rapport de suivi"
    synthese["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    synthese["A1"].fill = PatternFill("solid", fgColor=bleu_fonce)
    synthese["A1"].alignment = Alignment(horizontal="center")
    synthese["A3"] = "Indicateur"
    synthese["B3"] = "Valeur"
    for cellule in synthese[3]:
        cellule.fill = entete
        cellule.font = police_entete
        cellule.alignment = Alignment(horizontal="center")

    indicateurs = [
        ("Collaborateurs concernés", rapport.collaborateurs_concernes),
        ("Progressions totales", rapport.progressions_total),
        ("Formations terminées", rapport.progressions_terminees),
        ("Progressions en cours", rapport.progressions_en_cours),
        ("Formations en attente", rapport.formations_en_attente),
        ("Progression moyenne", rapport.taux_completion / 100),
        ("Taux de réussite", rapport.taux_reussite / 100 if rapport.taux_reussite is not None else None),
        ("Taux de conformité", rapport.taux_conformite / 100 if rapport.taux_conformite is not None else None),
    ]
    for ligne, (libelle, valeur) in enumerate(indicateurs, start=4):
        synthese.cell(ligne, 1, libelle)
        cellule_valeur = synthese.cell(ligne, 2, valeur if valeur is not None else "Non disponible")
        if ligne >= 9 and valeur is not None:
            cellule_valeur.number_format = "0%"
        if ligne % 2 == 0:
            synthese.cell(ligne, 1).fill = accent
            cellule_valeur.fill = accent
    synthese.column_dimensions["A"].width = 30
    synthese.column_dimensions["B"].width = 22

    colonnes = [
        "Collaborateur",
        "Service",
        "Formation",
        "Obligatoire",
        "Progression",
        "Statut",
        "Échéance",
        "Dernière activité",
    ]
    detail.append(colonnes)
    for cellule in detail[1]:
        cellule.fill = entete
        cellule.font = police_entete
        cellule.alignment = Alignment(horizontal="center")

    libelles_statut = {
        "non_commence": "En attente",
        "en_cours": "En cours",
        "termine": "Terminée",
    }
    for resultat in rapport.resultats:
        detail.append(
            [
                f"{resultat.utilisateur_prenom} {resultat.utilisateur_nom}",
                resultat.departement_nom or "Non renseigné",
                resultat.formation_titre,
                "Oui" if resultat.formation_obligatoire else "Non",
                resultat.pourcentage / 100,
                libelles_statut[resultat.statut],
                resultat.date_limite,
                resultat.derniere_activite,
            ]
        )

    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    for ligne in range(2, detail.max_row + 1):
        detail.cell(ligne, 5).number_format = "0%"
        detail.cell(ligne, 7).number_format = "dd/mm/yyyy"
        detail.cell(ligne, 8).number_format = "dd/mm/yyyy hh:mm"
    for colonne, largeur in enumerate([28, 22, 36, 14, 14, 16, 14, 22], start=1):
        detail.column_dimensions[get_column_letter(colonne)].width = largeur

    sortie = io.BytesIO()
    classeur.save(sortie)
    sortie.seek(0)
    nom_fichier = f"exia-academy-suivi-{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        sortie,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nom_fichier}"'},
    )


@router.get("/suivi/export/pdf")
def exporter_suivi_pdf(
    formation_id: uuid.UUID | None = Query(default=None),
    departement_id: uuid.UUID | None = Query(default=None),
    admin: Utilisateur = Depends(exiger_administrateur),
    session: Session = Depends(get_session),
) -> StreamingResponse:
    rapport = consulter_suivi(
        formation_id=formation_id,
        departement_id=departement_id,
        _=admin,
        session=session,
    )
    contenu_pdf = generer_pdf_rapport(rapport)
    nom_fichier = f"exia-academy-suivi-{date.today().isoformat()}.pdf"
    return StreamingResponse(
        io.BytesIO(contenu_pdf),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{nom_fichier}"'},
    )
